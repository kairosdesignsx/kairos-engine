"""Spreadsheet line: a budget workbook (.xlsx, opens in Google Sheets and Excel) built from a list of modules (tabs).

spec = {
  "title", "subtitle", "brand", "currency": "$", "year": 2027 | null (any year),
  "theme": {"header": "#hex", "accent": "#hex", "soft": "#hex", "ink": "#hex", "muted": "#hex", "paper": "#hex", "font": "Arial"},
  "categories": [{"name": "Rent", "group": "Needs", "budget": 1200}, ...],
  "income": [{"name": "Paycheck 1", "amount": 2400}, ...],
  "sample_data": true | false,
  "sections": [ {"module": "start_here", ...}, {"module": "setup"}, {"module": "months"}, {"module": "dashboard"}, ... ]
}
Every module registers itself in MODULES with a description and params (capabilities()).
Formula cells are locked and the sheets protected without a password, so buyers type only in input cells.
"""
import datetime, json, re
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter as COL
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

MONTHS = ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August', 'September', 'October', 'November', 'December']
GROUPS = ['Needs', 'Wants', 'Savings', 'Debt']
THEMES = {
    'sage':  dict(header='#2F3A2C', accent='#6E8B5E', soft='#EEF2EA', ink='#22281F', muted='#6B7466', paper='#FFFFFF', warn='#C8553D', font='Arial'),
    'slate': dict(header='#1F2D3A', accent='#2C7A74', soft='#E6F0EC', ink='#1F2D3A', muted='#5B6874', paper='#FFFFFF', warn='#E07A5F', font='Arial'),
    'blush': dict(header='#4A3036', accent='#C98D94', soft='#FBEDEC', ink='#3A2A2E', muted='#8A6F74', paper='#FFFFFF', warn='#B5473A', font='Arial'),
    'night': dict(header='#ECE8F5', accent='#8C7BD8', soft='#2A2F42', ink='#ECE8F5', muted='#A3A0B8', paper='#1E2230', warn='#F3B79B', font='Arial'),
}
MODULES = {}
EXTRAS = ('bills', 'paycheck', 'debt', 'savings')  # optional tabs: muted tab colour so the core path stands out


def module(name, desc, params=None):
    def deco(fn):
        MODULES[name] = {'fn': fn, 'desc': desc, 'params': params or {}}
        return fn
    return deco


def esc(s):  # photos.py shares this helper name with the planner line
    import html
    return html.escape(str(s))


def resolve_theme(t):
    th = dict(THEMES[t] if isinstance(t, str) else THEMES[(t or {}).get('base', 'slate')])
    if isinstance(t, dict):
        th.update({k: v for k, v in t.items() if k != 'base'})
    # photos.py expects these names
    th.setdefault('display', th.get('font', 'Arial')); th.setdefault('body', th.get('font', 'Arial'))
    th.setdefault('line', th['soft']); th.setdefault('warm', th['warn'])
    return th


def motif(th, where, seed=1, size=150):
    return ''


def hexc(c):
    return c.lstrip('#').upper()


class Book:
    def __init__(self, spec):
        self.s = spec
        self.th = resolve_theme(spec.get('theme', 'slate'))
        self.cur = spec.get('currency', '$')
        self.money = f'"{self.cur}"#,##0.00;[Red]-"{self.cur}"#,##0.00;"{self.cur}"0.00'
        self.year = spec.get('year')
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.sheets = []        # (anchor, title, module)
        self.nav = []           # (label, sheet title) for the link bar on each tab
        self.cats = spec.get('categories') or default_categories()
        self.income = spec.get('income') or [{'name': 'Paycheck', 'amount': 3000}, {'name': 'Side income', 'amount': 0}]
        self.month_titles = [m[:3] for m in MONTHS]
        self.current = None
        self.after = []         # callbacks run once every tab exists (for tabs that read later tabs)

    # ---- styles ----
    def font(self, size=11, bold=False, color=None, italic=False):
        return Font(name=self.th['font'], size=size, bold=bold, italic=italic, color=hexc(color or self.th['ink']))

    def fill(self, c):
        return PatternFill('solid', fgColor=hexc(c))

    def sheet(self, anchor, title, tab_color=None):
        ws = self.wb.create_sheet(title[:31])
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = hexc(tab_color or (self.th['muted'] if self.current in EXTRAS else self.th['accent']))
        for c in range(1, 30):
            ws.column_dimensions[COL(c)].width = 14
        ws.column_dimensions['A'].width = 3
        if self.th['paper'].upper() != '#FFFFFF':
            for r in range(1, 140):
                for c in range(1, 22):
                    ws.cell(r, c).fill = self.fill(self.th['paper'])
        # one landscape page per tab when printed or exported (also how the engine takes its screenshots)
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.gridLines = False
        self.sheets.append((anchor, ws.title, self.current))
        return ws

    def title(self, ws, text, sub=''):
        ws.row_dimensions[2].height = 34
        c = ws.cell(2, 2, text); c.font = self.font(22, True)
        if sub:
            s = ws.cell(3, 2, sub); s.font = self.font(11, color=self.th['muted'])
        ws.cell(4, 2).value = None  # nav bar row filled at the end

    def head(self, ws, r, c, labels, widths=None):
        for i, l in enumerate(labels):
            cell = ws.cell(r, c + i, l)
            cell.font = self.font(10, True, color='#FFFFFF' if self.th['paper'].upper() == '#FFFFFF' else self.th['paper'])
            cell.fill = self.fill(self.th['header'])
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            if widths and i < len(widths) and widths[i]:
                ws.column_dimensions[COL(c + i)].width = widths[i]
        ws.row_dimensions[r].height = 22

    def inp(self, cell, value=None, fmt=None):
        """An input cell: unlocked, soft fill, so buyers know where to type."""
        if value is not None:
            cell.value = value
        cell.protection = Protection(locked=False)
        cell.fill = self.fill(self.th['soft'])
        cell.font = self.font(11)
        cell.border = Border(bottom=Side(style='thin', color=hexc(self.th['muted'])))
        if fmt:
            cell.number_format = fmt

    def calc(self, cell, formula, fmt=None, bold=False):
        cell.value = formula
        cell.font = self.font(11, bold)
        if fmt:
            cell.number_format = fmt

    def protect(self, ws):
        ws.protection.sheet = True
        ws.protection.formatColumns = False
        ws.protection.formatRows = False
        ws.protection.insertRows = False
        ws.protection.sort = False
        ws.protection.autoFilter = False

    def save(self, path):
        # nav bar: a link to every main tab on row 4 of each sheet
        for anchor, title, mod in self.sheets:
            ws = self.wb[title]
            c0 = 2
            while (ws.column_dimensions[COL(c0)].width or 14) < 10:
                c0 += 1
            for i, (label, target) in enumerate(self.nav[:9]):
                c = ws.cell(4, c0 + i, label)
                c.hyperlink = f"#'{target}'!A1"
                c.font = self.font(9, True, color=self.th['accent'])
                c.alignment = Alignment(horizontal='left')
        from openpyxl.workbook.properties import CalcProperties
        self.wb.calculation = CalcProperties(fullCalcOnLoad=True)  # Excel recalculates on open; Google Sheets always does
        self.wb.save(path)


def default_categories():
    return [{'name': n, 'group': g, 'budget': b} for n, g, b in [
        ('Rent or mortgage', 'Needs', 1200), ('Utilities', 'Needs', 180), ('Groceries', 'Needs', 450), ('Transport', 'Needs', 200),
        ('Insurance', 'Needs', 150), ('Phone and internet', 'Needs', 90), ('Eating out', 'Wants', 150), ('Shopping', 'Wants', 120),
        ('Subscriptions', 'Wants', 45), ('Fun', 'Wants', 100), ('Emergency fund', 'Savings', 200), ('Sinking funds', 'Savings', 100),
        ('Credit card', 'Debt', 150), ('Student loan', 'Debt', 120)]]


# ---------------- modules ----------------
@module('start_here', 'Start here tab: numbered steps, how to copy into Google Sheets or open in Excel, what the colors mean, links to every tab. steps: [[title, text]], tip',
        params={'steps': 'list', 'tip': 'str'})
def m_start(B, sec):
    ws = B.sheet('start', sec.get('label', 'Start here'))
    B.title(ws, B.s['title'], B.s.get('subtitle', ''))
    ws.column_dimensions['B'].width = 6; ws.column_dimensions['C'].width = 36; ws.column_dimensions['D'].width = 70
    r = 6
    for i, st in enumerate(sec.get('steps') or [], 1):
        n = ws.cell(r, 2, i); n.font = B.font(14, True, '#FFFFFF'); n.fill = B.fill(B.th['accent']); n.alignment = Alignment(horizontal='center', vertical='center')
        t = ws.cell(r, 3, st[0]); t.font = B.font(13, True); t.alignment = Alignment(vertical='center', wrap_text=True)
        x = ws.cell(r, 4, st[1]); x.font = B.font(11, color=B.th['muted']); x.alignment = Alignment(vertical='center', wrap_text=True)
        ws.row_dimensions[r].height = 48
        r += 1
    r += 1
    ws.cell(r, 3, 'What the colors mean').font = B.font(12, True); r += 1
    a = ws.cell(r, 3, 'Type here'); B.inp(a); ws.cell(r, 4, 'Shaded cells are yours to fill in.').font = B.font(11, color=B.th['muted']); r += 1
    b = ws.cell(r, 3, 'Calculated'); b.font = B.font(11, True); ws.cell(r, 4, 'Plain cells do the math. They are locked so formulas cannot break.').font = B.font(11, color=B.th['muted']); r += 2
    if sec.get('tip'):
        tcell = ws.cell(r, 3, 'Tip'); tcell.font = B.font(12, True)
        tx = ws.cell(r, 4, sec['tip']); tx.font = B.font(11); tx.alignment = Alignment(wrap_text=True); ws.row_dimensions[r].height = 40
    B.nav.insert(0, ('Start here', ws.title))
    B.protect(ws)


@module('setup', 'Setup tab: year, currency note, income sources and spending categories with group (Needs, Wants, Savings, Debt) and monthly budget. Every other tab reads from here. rows: number of category rows (default 30)',
        params={'rows': 'int'})
def m_setup(B, sec):
    ws = B.sheet('setup', sec.get('label', 'Setup'))
    B.title(ws, 'Setup', 'Fill this in once. Every other tab updates from here.')
    ws.cell(6, 2, 'Year').font = B.font(11, True); B.inp(ws.cell(6, 3), B.year or datetime.date.today().year + 1, '0')
    ws.cell(7, 2, 'Currency').font = B.font(11, True); ws.cell(7, 3, B.cur).font = B.font(11)
    B.wb.defined_names['Year'] = DefinedName('Year', attr_text=f"'{ws.title}'!$C$6")
    # income
    B.head(ws, 9, 2, ['Income source', 'Monthly amount'], [28, 16])
    for i in range(8):
        it = B.income[i] if i < len(B.income) else {'name': None, 'amount': None}
        B.inp(ws.cell(10 + i, 2), it['name']); B.inp(ws.cell(10 + i, 3), it['amount'], B.money)
    ws.cell(18, 2, 'Expected income').font = B.font(11, True)
    B.calc(ws.cell(18, 3), '=SUM(C10:C17)', B.money, True)
    # categories
    n = sec.get('rows', 30)
    B.head(ws, 9, 5, ['Category', 'Group', 'Monthly budget'], [28, 14, 16])
    dv = DataValidation(type='list', formula1='"' + ','.join(GROUPS) + '"', allow_blank=True); ws.add_data_validation(dv)
    for i in range(n):
        c = B.cats[i] if i < len(B.cats) else {'name': None, 'group': None, 'budget': None}
        B.inp(ws.cell(10 + i, 5), c['name']); g = ws.cell(10 + i, 6); B.inp(g, c['group']); dv.add(g)
        B.inp(ws.cell(10 + i, 7), c['budget'], B.money)
    last = 9 + n
    ws.cell(last + 1, 5, 'Total budget').font = B.font(11, True)
    B.calc(ws.cell(last + 1, 7), f'=SUM(G10:G{last})', B.money, True)
    ws.cell(last + 2, 5, 'Left to budget').font = B.font(11, True)
    B.calc(ws.cell(last + 2, 7), f'=C18-G{last + 1}', B.money, True)
    ws.conditional_formatting.add(f'G{last + 2}', CellIsRule(operator='lessThan', formula=['0'], font=Font(color=hexc(B.th['warn']), bold=True)))
    B.wb.defined_names['Categories'] = DefinedName('Categories', attr_text=f"'{ws.title}'!$E$10:$E${last}")
    B.setup = {'title': ws.title, 'cat_first': 10, 'cat_last': last, 'income_total': f"'{ws.title}'!$C$18"}
    B.nav.append(('Setup', ws.title))
    ws.freeze_panes = 'A9'
    B.protect(ws)


@module('months', 'Twelve month tabs: income received, a spending log with a category dropdown, and budget vs actual per category with progress and over-budget warnings. log_rows (default 80)',
        params={'log_rows': 'int'})
def m_months(B, sec):
    S = B.setup; n_log = sec.get('log_rows', 80)
    ncat = S['cat_last'] - S['cat_first'] + 1
    B.month_sheets = []
    for mi, mname in enumerate(MONTHS):
        ws = B.sheet(f'm{mi + 1:02d}', mname)
        B.title(ws, mname, 'Log what you spend. The summary on the right updates itself.')
        # spending log
        B.head(ws, 6, 2, ['Date', 'What', 'Category', 'Amount'], [12, 26, 22, 14])
        dv = DataValidation(type='list', formula1='=Categories', allow_blank=True); ws.add_data_validation(dv)
        for r in range(7, 7 + n_log):
            B.inp(ws.cell(r, 2), None, 'mmm d'); B.inp(ws.cell(r, 3)); c = ws.cell(r, 4); B.inp(c); dv.add(c); B.inp(ws.cell(r, 5), None, B.money)
        logl = 6 + n_log
        # month totals at the top right, so the answer to "how am I doing" is the first thing on the tab
        t = 6
        rows = [('Income this month', f"={S['income_total']}"), ('Spent', f'=SUM(E7:E{logl})'), ('Left over', f'=H{t}-H{t + 1}'),
                ('Saved and debt paid', f"=SUMPRODUCT(SUMIFS($E$7:$E${logl},$D$7:$D${logl},'{S['title']}'!$E${S['cat_first']}:$E${S['cat_last']})*(('{S['title']}'!$F${S['cat_first']}:$F${S['cat_last']}=\"Savings\")+('{S['title']}'!$F${S['cat_first']}:$F${S['cat_last']}=\"Debt\")))")]
        for k, (lab, f) in enumerate(rows):
            lc = ws.cell(t + k, 7, lab); lc.font = B.font(11, True); lc.fill = B.fill(B.th['soft'] if k != 0 else B.th['paper'])
            if k == 0:  # income: starts as the expected income from Setup, buyers can type over it
                B.inp(ws.cell(t, 8), f, B.money); ws.cell(t, 8).font = B.font(12, True)
                ws.cell(t, 9, 'Type over this if your income was different').font = B.font(9, italic=True, color=B.th['muted'])
            else:
                B.calc(ws.cell(t + k, 8), f, B.money, True); ws.cell(t + k, 8).font = B.font(12, True)
        ws.conditional_formatting.add(f'H{t + 2}', CellIsRule(operator='lessThan', formula=['0'], font=Font(color=hexc(B.th['warn']), bold=True)))
        # budget vs actual by category
        h = 11
        B.head(ws, h, 7, ['Category', 'Budget', 'Spent', 'Left', 'Used'], [22, 13, 13, 13, 10])
        for i in range(ncat):
            r = h + 1 + i; sr = S['cat_first'] + i
            B.calc(ws.cell(r, 7), f"=IF('{S['title']}'!E{sr}=\"\",\"\",'{S['title']}'!E{sr})")
            B.calc(ws.cell(r, 8), f"=IF(G{r}=\"\",\"\",'{S['title']}'!G{sr})", B.money)
            B.calc(ws.cell(r, 9), f'=IF(G{r}="","",SUMIFS($E$7:$E${logl},$D$7:$D${logl},G{r}))', B.money)
            B.calc(ws.cell(r, 10), f'=IF(G{r}="","",H{r}-I{r})', B.money)
            B.calc(ws.cell(r, 11), f'=IF(OR(G{r}="",H{r}=0),"",I{r}/H{r})', '0%')
        last = h + ncat
        ws.conditional_formatting.add(f'J{h + 1}:J{last}', CellIsRule(operator='lessThan', formula=['0'], font=Font(color=hexc(B.th['warn']), bold=True)))
        ws.conditional_formatting.add(f'K{h + 1}:K{last}', DataBarRule(start_type='num', start_value=0, end_type='num', end_value=1, color=hexc(B.th['accent'])))
        ws.freeze_panes = 'A7'
        B.month_sheets.append({'title': ws.title, 'income': f'H{t}', 'spent': f'H{t + 1}', 'left': f'H{t + 2}', 'saved': f'H{t + 3}',
                               'log': (7, logl)})
        B.protect(ws)
    B.nav.append(('Months', B.month_sheets[0]['title']))


@module('dashboard', 'Dashboard tab: the year by month (income, spent, left over, saved) and spending by category for the year, with a bar chart and a pie chart. Needs months.')
def m_dashboard(B, sec):
    S = B.setup; ws = B.sheet('dashboard', sec.get('label', 'Dashboard'))
    B.title(ws, sec.get('label', 'Dashboard'), 'Your year at a glance. Nothing to type here.')
    B.head(ws, 6, 2, ['Month', 'Income', 'Spent', 'Left over', 'Saved'], [14, 14, 14, 14, 14])
    for i, M in enumerate(B.month_sheets):
        r = 7 + i
        ws.cell(r, 2, MONTHS[i]).hyperlink = f"#'{M['title']}'!A1"; ws.cell(r, 2).font = B.font(11, color=B.th['accent'])
        active = f"COUNT('{M['title']}'!$E${M['log'][0]}:$E${M['log'][1]})>0"
        for j, key in enumerate(['income', 'spent', 'left', 'saved']):
            # months with nothing logged yet stay empty so the year totals only count months in use
            B.calc(ws.cell(r, 3 + j), f"=IF({active},'{M['title']}'!{M[key]},\"\")", B.money)
    ws.cell(19, 2, 'Year').font = B.font(11, True)
    for j in range(4):
        B.calc(ws.cell(19, 3 + j), f'=SUM({COL(3 + j)}7:{COL(3 + j)}18)', B.money, True)
    ws.cell(20, 2, 'Savings rate').font = B.font(11, True)
    B.calc(ws.cell(20, 3), '=IF(C19=0,"",F19/C19)', '0%', True)
    # category totals for the year
    B.head(ws, 6, 8, ['Category', 'Spent this year'], [24, 16])
    ncat = S['cat_last'] - S['cat_first'] + 1
    for i in range(ncat):
        r = 7 + i; sr = S['cat_first'] + i
        B.calc(ws.cell(r, 8), f"=IF('{S['title']}'!E{sr}=\"\",\"\",'{S['title']}'!E{sr})")
        parts = '+'.join(f"SUMIFS('{M['title']}'!$E${M['log'][0]}:$E${M['log'][1]},'{M['title']}'!$D${M['log'][0]}:$D${M['log'][1]},H{r})" for M in B.month_sheets)
        B.calc(ws.cell(r, 9), f'=IF(H{r}="","",{parts})', B.money)
    # spending by group (Needs, Wants, Savings, Debt): always complete, however many categories the buyer adds
    for i in range(ncat):
        r = 7 + i; sr = S['cat_first'] + i
        B.calc(ws.cell(r, 10), f"=IF(H{r}=\"\",\"\",'{S['title']}'!F{sr})")
    ws.column_dimensions['J'].hidden = True
    g0 = 23
    B.head(ws, g0, 2, ['Group', 'Spent this year'])
    for k, grp in enumerate(GROUPS):
        r = g0 + 1 + k
        ws.cell(r, 2, grp).font = B.font(11)
        B.calc(ws.cell(r, 3), f'=SUMIF($J$7:$J${6 + ncat},"{grp}",$I$7:$I${6 + ncat})', B.money)
    top = g0 + 14  # bar chart under the group table and pie; the pie sits beside the group table
    bar = BarChart(); bar.type = 'col'; bar.title = 'Income and spending by month'; bar.height = 7.5; bar.width = 16
    bar.add_data(Reference(ws, min_col=3, max_col=4, min_row=6, max_row=18), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=2, min_row=7, max_row=18))
    for ser, col in zip(bar.series, [B.th['accent'], B.th['header']]):
        ser.graphicalProperties.solidFill = hexc(col); ser.graphicalProperties.line.solidFill = hexc(col)
    ws.add_chart(bar, f'B{top}')
    pie = PieChart(); pie.title = 'Where the money went'; pie.height = 7.5; pie.width = 11
    pie.add_data(Reference(ws, min_col=3, min_row=g0, max_row=g0 + 4), titles_from_data=True)
    pie.set_categories(Reference(ws, min_col=2, min_row=g0 + 1, max_row=g0 + 4))
    pie.dataLabels = DataLabelList(); pie.dataLabels.showPercent = True
    pie.dataLabels.showVal = False; pie.dataLabels.showCatName = False; pie.dataLabels.showSerName = False; pie.dataLabels.showLeaderLines = False
    from openpyxl.chart.series import DataPoint
    for k, col in enumerate([B.th['header'], B.th['accent'], '#8FB8A8', B.th['warn']]):
        pt = DataPoint(idx=k); pt.graphicalProperties.solidFill = hexc(col); pie.series[0].dPt.append(pt)
    ws.add_chart(pie, f'D{g0 - 1}')
    B.nav.insert(1 if B.nav and B.nav[0][0] == 'Start here' else 0, ('Dashboard', ws.title))
    B.protect(ws)


@module('bills', 'Bill tracker: each bill with due day, amount, autopay, and a paid mark for every month. rows (default 25)', params={'rows': 'int'})
def m_bills(B, sec):
    ws = B.sheet('bills', sec.get('label', 'Bills'))
    B.title(ws, sec.get('label', 'Bills'), 'Mark each bill paid with a ✓ from the dropdown.')
    B.head(ws, 6, 2, ['Bill', 'Due day', 'Amount', 'Autopay'] + [m[:3] for m in MONTHS], [24, 9, 12, 9] + [6] * 12)
    dv = DataValidation(type='list', formula1='"✓"', allow_blank=True); ws.add_data_validation(dv)
    yn = DataValidation(type='list', formula1='"Yes,No"', allow_blank=True); ws.add_data_validation(yn)
    for r in range(7, 7 + sec.get('rows', 25)):
        B.inp(ws.cell(r, 2)); B.inp(ws.cell(r, 3), None, '0'); B.inp(ws.cell(r, 4), None, B.money); a = ws.cell(r, 5); B.inp(a); yn.add(a)
        for m in range(12):
            c = ws.cell(r, 6 + m); B.inp(c); dv.add(c); c.alignment = Alignment(horizontal='center')
    last = 6 + sec.get('rows', 25)
    B.bills = {'title': ws.title, 'first': 7, 'last': last}
    ws.cell(last + 1, 2, 'Total per month').font = B.font(11, True)
    B.calc(ws.cell(last + 1, 4), f'=SUM(D7:D{last})', B.money, True)
    B.nav.append(('Bills', ws.title))
    ws.freeze_panes = 'A7'
    B.protect(ws)


@module('debt', 'Debt payoff: each debt with balance, interest rate and minimum; choose Snowball or Avalanche and an extra monthly amount; shows payoff order and months to pay off at the minimum. rows (default 10)', params={'rows': 'int'})
def m_debt(B, sec):
    ws = B.sheet('debt', sec.get('label', 'Debt payoff'))
    B.title(ws, sec.get('label', 'Debt payoff'), 'List your debts. The order and timelines update as you pay.')
    ws.cell(6, 2, 'Method').font = B.font(11, True)
    m = ws.cell(6, 3); B.inp(m, 'Snowball'); dv = DataValidation(type='list', formula1='"Snowball,Avalanche"'); ws.add_data_validation(dv); dv.add(m)
    ws.cell(7, 2, 'Extra each month').font = B.font(11, True); B.inp(ws.cell(7, 3), 100, B.money)
    B.head(ws, 9, 2, ['Debt', 'Balance', 'Rate (APR)', 'Minimum', 'Pay order', 'Months at minimum', 'Interest at minimum'], [24, 14, 12, 12, 11, 16, 18])
    n = sec.get('rows', 10)
    for r in range(10, 10 + n):
        B.inp(ws.cell(r, 2)); B.inp(ws.cell(r, 3), None, B.money); B.inp(ws.cell(r, 4), None, '0.0%'); B.inp(ws.cell(r, 5), None, B.money)
        last = 9 + n
        B.calc(ws.cell(r, 6), f'=IF(C{r}="","",IF($C$6="Snowball",COUNTIFS($C$10:$C${last},"<"&C{r},$C$10:$C${last},"<>")+1,COUNTIFS($D$10:$D${last},">"&D{r},$C$10:$C${last},"<>")+1))', '0')
        B.calc(ws.cell(r, 7), f'=IF(OR(C{r}="",E{r}=""),"",IFERROR(ROUNDUP(NPER(D{r}/12,-E{r},C{r}),0),"never: raise the minimum"))', '0')
        B.calc(ws.cell(r, 8), f'=IF(OR(C{r}="",E{r}="",NOT(ISNUMBER(G{r}))),"",G{r}*E{r}-C{r})', B.money)
    last = 9 + n
    ws.cell(last + 1, 2, 'Total').font = B.font(11, True)
    B.calc(ws.cell(last + 1, 3), f'=SUM(C10:C{last})', B.money, True); B.calc(ws.cell(last + 1, 5), f'=SUM(E10:E{last})', B.money, True)
    B.nav.append(('Debt', ws.title))
    B.protect(ws)


@module('savings', 'Savings goals: goal, target, saved so far, progress bar, target date and how much to save per month. rows (default 10)', params={'rows': 'int'})
def m_savings(B, sec):
    ws = B.sheet('savings', sec.get('label', 'Savings goals'))
    B.title(ws, sec.get('label', 'Savings goals'), 'One line per goal. The bar fills as you save.')
    B.head(ws, 6, 2, ['Goal', 'Target', 'Saved', 'Progress', 'Target date', 'Save per month'], [24, 14, 14, 12, 14, 16])
    n = sec.get('rows', 10)
    for r in range(7, 7 + n):
        B.inp(ws.cell(r, 2)); B.inp(ws.cell(r, 3), None, B.money); B.inp(ws.cell(r, 4), None, B.money)
        B.calc(ws.cell(r, 5), f'=IF(OR(C{r}="",C{r}=0),"",MIN(1,D{r}/C{r}))', '0%')
        B.inp(ws.cell(r, 6), None, 'mmm yyyy')
        B.calc(ws.cell(r, 7), f'=IF(OR(C{r}="",F{r}=""),"",MAX(0,(C{r}-D{r})/MAX(1,DATEDIF(TODAY(),F{r},"m"))))', B.money)
    ws.conditional_formatting.add(f'E7:E{6 + n}', DataBarRule(start_type='num', start_value=0, end_type='num', end_value=1, color=hexc(B.th['accent'])))
    B.nav.append(('Savings', ws.title))
    B.protect(ws)


@module('start_check', 'Start here tab with a live Ready check: steps, colour key, tip, then one row per check (income, balanced, groups, orphans, blanks, example, split) showing Done or To do and what to fix, a jump row to all 12 months, and an If something looks wrong list. Needs setup and months. Anchor start.',
        params={'label': 'str', 'steps': 'list of [title, text]', 'tip': 'str', 'checks': 'list of {id, label, fix}', 'month_jump': 'bool', 'if_wrong': 'list of str'})
def m_start_check(B, sec):
    ws = B.sheet('start', sec.get('label', 'Start here'))
    B.title(ws, B.s['title'], B.s.get('subtitle', ''))
    ws.column_dimensions['B'].width = 6; ws.column_dimensions['C'].width = 34; ws.column_dimensions['D'].width = 14; ws.column_dimensions['E'].width = 60
    r = 6
    for i, st in enumerate(sec.get('steps') or [], 1):
        n = ws.cell(r, 2, i); n.font = B.font(14, True, '#FFFFFF'); n.fill = B.fill(B.th['accent']); n.alignment = Alignment(horizontal='center', vertical='center')
        t = ws.cell(r, 3, st[0]); t.font = B.font(12, True); t.alignment = Alignment(vertical='center', wrap_text=True)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
        x = ws.cell(r, 4, st[1]); x.font = B.font(10, color=B.th['muted']); x.alignment = Alignment(vertical='center', wrap_text=True)
        ws.row_dimensions[r].height = 15 * max(2, len(st[1]) // 70 + 1)
        r += 1
    r += 1
    ready = r
    ws.cell(r, 3, 'Ready check').font = B.font(14, True); ws.cell(r, 5, 'Updates as you type.').font = B.font(10, italic=True, color=B.th['muted']); r += 1
    B.head(ws, r, 3, ['Check', 'Status', 'What to do']); r += 1
    rows = {}
    for c in sec.get('checks') or []:
        if c['id'] == 'example' and not B.s.get('sample_data'):
            continue
        ws.cell(r, 3, c['label']).font = B.font(11)
        ws.cell(r, 5).alignment = Alignment(wrap_text=True, vertical='center')
        rows[c['id']] = (r, c); r += 1
    first_check, last_check = ready + 2, r - 1
    r += 1
    jump = r
    if sec.get('month_jump', True):
        ws.cell(r, 3, 'Jump to a month').font = B.font(12, True); r += 1
        r += 5  # 4 rows of 3 month links (filled once the month tabs exist) and a gap
    wrong_r = r
    if sec.get('if_wrong'):
        ws.cell(r, 3, 'If something looks wrong').font = B.font(12, True); r += 1
        for line in sec['if_wrong']:
            ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
            x = ws.cell(r, 3, line); x.font = B.font(10); x.alignment = Alignment(wrap_text=True, vertical='center')
            ws.row_dimensions[r].height = 15 * max(1, len(line) // 95 + 1)
            r += 1
    r += 1
    ws.cell(r, 3, 'Colour key').font = B.font(12, True); r += 1
    a = ws.cell(r, 3, 'Type here'); B.inp(a); ws.cell(r, 5, 'Shaded cells are yours to fill in.').font = B.font(10, color=B.th['muted']); r += 1
    b = ws.cell(r, 3, 'Calculated'); b.font = B.font(11, True); ws.cell(r, 5, 'Plain cells hold formulas.').font = B.font(10, color=B.th['muted']); r += 1
    if sec.get('tip'):
        r += 1
        tl = ws.cell(r, 3, 'Tip'); tl.font = B.font(12, True); tl.alignment = Alignment(vertical='center')
        tx = ws.cell(r, 4, sec['tip']); ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
        tx.font = B.font(10); tx.alignment = Alignment(wrap_text=True, vertical='center'); ws.row_dimensions[r].height = 15 * (len(sec['tip']) // 70 + 1)
    B.nav.insert(0, ('Start here', ws.title))
    B.protect(ws)

    def fill():
        S = B.setup; st = S['title']; f0, f1 = S['cat_first'], S['cat_last']
        L = f"'{st}'!$G${f1 + 2}"; T = f"'{st}'!$G${f1 + 1}"
        money = f'"{B.cur}#,##0.00"'
        months = getattr(B, 'month_sheets', [])
        # helper columns (hidden): per-month counts for orphans and blanks
        hc = 20
        for i, M in enumerate(months):
            a0, a1 = M['log']; t = M['title']
            rr = 6 + i
            ws.cell(rr, hc, M['title'])
            ws.cell(rr, hc + 1, f"=SUMPRODUCT(('{t}'!$E${a0}:$E${a1}<>\"\")*('{t}'!$D${a0}:$D${a1}<>\"\")*(COUNTIF(Categories,'{t}'!$D${a0}:$D${a1})=0))")
            ws.cell(rr, hc + 2, f"=SUMPRODUCT(('{t}'!$E${a0}:$E${a1}<>\"\")*('{t}'!$D${a0}:$D${a1}=\"\"))")
            ws.cell(rr, hc + 3, f'=IF({COL(hc + 1)}{rr}>0,1,0)')
        for k in range(4):
            ws.column_dimensions[COL(hc + k)].hidden = True
        h0, h1 = 6, 6 + max(0, len(months) - 1)
        HC, HO, HB, HF = COL(hc), COL(hc + 1), COL(hc + 2), COL(hc + 3)
        F = {}
        F['income'] = (f"=IF({S['income_total']}>0,\"Done\",\"To do\")", f"=IF({S['income_total']}>0,\"\",\"{_q(rows.get('income', (0, {'fix': ''}))[1].get('fix') or 'Add your take-home pay in Setup')}\")")
        F['balanced'] = (f"=IF(ROUND({L},2)=0,\"Done\",\"To do\")",
                         f"=IF(ROUND({L},2)>0,\"You have \"&TEXT({L},{money})&\" not budgeted yet\",IF(ROUND({L},2)<0,\"Over by \"&TEXT(-{L},{money})&\": lower a budget\",\"\"))")
        F['groups'] = (f"=IF(COUNTIFS('{st}'!$E${f0}:$E${f1},\"<>\",'{st}'!$F${f0}:$F${f1},\"\")=0,\"Done\",\"To do\")",
                       f"=IF(COUNTIFS('{st}'!$E${f0}:$E${f1},\"<>\",'{st}'!$F${f0}:$F${f1},\"\")=0,\"\",COUNTIFS('{st}'!$E${f0}:$E${f1},\"<>\",'{st}'!$F${f0}:$F${f1},\"\")&\" categories need a group: pick Needs, Wants, Savings or Debt in Setup\")")
        F['orphans'] = (f"=IF(SUM({HO}{h0}:{HO}{h1})=0,\"Done\",\"To do\")",
                        f"=IF(SUM({HO}{h0}:{HO}{h1})=0,\"\",SUM({HO}{h0}:{HO}{h1})&\" entries use a category not in Setup (first in \"&INDEX({HC}{h0}:{HC}{h1},MATCH(1,{HF}{h0}:{HF}{h1},0))&\"): pick their category again\")")
        F['blanks'] = (f"=IF(SUM({HB}{h0}:{HB}{h1})=0,\"Done\",\"To do\")",
                       f"=IF(SUM({HB}{h0}:{HB}{h1})=0,\"\",SUM({HB}{h0}:{HB}{h1})&\" amounts have no category, so they are not counted yet\")")
        if months and B.s.get('sample_data'):
            j = months[0]; a0, a1 = j['log']
            texts = ','.join(f'"{_q(t)}"' for _, t, c, _ in SAMPLE_ROWS)
            cnt = f"SUMPRODUCT(COUNTIF('{j['title']}'!$C${a0}:$C${a1},{{{texts}}}))"
            F['example'] = (f'=IF({cnt}=0,"Done","To do")', f'=IF({cnt}=0,"","Clear the example in {j["title"]} (step 2)")')
        g = lambda grp: f"SUMIF('{st}'!$F${f0}:$F${f1},\"{grp}\",'{st}'!$G${f0}:$G${f1})"
        F['split'] = ('Info', f"=IF({T}=0,\"\",TEXT({g('Needs')}/{T},\"0%\")&\" needs, \"&TEXT({g('Wants')}/{T},\"0%\")&\" wants, \"&TEXT(({g('Savings')}+{g('Debt')})/{T},\"0%\")&\" savings and debt (50/30/20 guide)\")")
        for cid, (rr, c) in rows.items():
            if cid not in F:
                continue
            stt, fix = F[cid]
            sc = ws.cell(rr, 4, stt); sc.font = B.font(11, True); sc.alignment = Alignment(horizontal='center')
            ws.cell(rr, 5, fix).font = B.font(10)
        rng = f'D{first_check}:D{last_check}'
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'D{first_check}="Done"'], fill=B.fill(B.th['accent']), font=Font(color='FFFFFF', bold=True)))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'D{first_check}="To do"'], font=Font(color=hexc(B.th['warn']), bold=True)))
        if sec.get('month_jump', True):
            for i, M in enumerate(months):
                c = ws.cell(jump + 1 + i // 3, 3 + (i % 3) if i % 3 < 2 else 5, M['title'])
                c.hyperlink = f"#'{M['title']}'!A1"; c.font = B.font(11, True, color=B.th['accent'])
    B.after.append(fill)


def _q(s):
    return str(s).replace('"', '""')


@module('paycheck', 'Paycheck plan tab: next payday, how often (Weekly, Every 2 weeks, Twice a month, Monthly) and usual take-home; one row per payday to the end of the year with the bills due before the next payday and what is left. Needs setup and bills. Anchor paycheck.',
        params={'label': 'str', 'rows': 'int (default 27)', 'frequency': 'str'})
def m_paycheck(B, sec):
    S = B.setup
    if not getattr(B, 'bills', None):
        raise ValueError('paycheck needs a bills section before it')
    ws = B.sheet('paycheck', sec.get('label', 'Paycheck plan'))
    B.title(ws, sec.get('label', 'Paycheck plan'), 'Paid weekly or every two weeks? See which bills each paycheck covers. Paid monthly? Skip this tab.')
    FREQ = ['Weekly', 'Every 2 weeks', 'Twice a month', 'Monthly']
    ws.cell(6, 2, 'Next payday').font = B.font(11, True)
    B.inp(ws.cell(6, 3), '=DATE(Year,1,1)+MOD(6-WEEKDAY(DATE(Year,1,1)),7)', 'mmm d, yyyy')
    ws.cell(7, 2, 'How often').font = B.font(11, True)
    fq = ws.cell(7, 3); B.inp(fq, sec.get('frequency', 'Every 2 weeks'))
    dv = DataValidation(type='list', formula1='"' + ','.join(FREQ) + '"'); ws.add_data_validation(dv); dv.add(fq)
    ws.cell(8, 2, 'Usual paycheck').font = B.font(11, True)
    B.inp(ws.cell(8, 3), f"=ROUND('{S['title']}'!$C$10*12/CHOOSE(MATCH($C$7,{{\"Weekly\",\"Every 2 weeks\",\"Twice a month\",\"Monthly\"}},0),52,26,24,12),2)", B.money)
    ws.cell(8, 4, 'Starts from your first income line in Setup. Type your real take-home pay.').font = B.font(9, italic=True, color=B.th['muted'])
    B.head(ws, 10, 2, ['Payday', 'Paycheck', 'Bills due before next payday', 'Left from this paycheck', 'Next payday'], [16, 14, 18, 18, 14])
    bl = B.bills
    n = sec.get('rows', 27)
    idx = 'MATCH($C$7,{"Weekly","Every 2 weeks","Twice a month","Monthly"},0)'
    nxt = lambda p: f'CHOOSE({idx},{p}+7,{p}+14,IF(DAY({p})<15,DATE(YEAR({p}),MONTH({p}),15),DATE(YEAR({p}),MONTH({p})+1,1)),EDATE({p},1))'
    for i in range(n):
        r = 11 + i
        if i == 0:
            B.calc(ws.cell(r, 2), '=$C$6', 'ddd mmm d')
        else:
            B.calc(ws.cell(r, 2), f'=IF(F{r - 1}="","",IF(F{r - 1}>DATE(Year,12,31),"",F{r - 1}))', 'ddd mmm d')
        B.inp(ws.cell(r, 3), f'=IF(B{r}="","",$C$8)', B.money)
        B.calc(ws.cell(r, 6), f'=IF(B{r}="","",{nxt(f"B{r}")})', 'mmm d')
        due = f"'{bl['title']}'!$C${bl['first']}:$C${bl['last']}"
        amt = f"'{bl['title']}'!$D${bl['first']}:$D${bl['last']}"
        B.calc(ws.cell(r, 4), f'=IF(B{r}="","",SUMPRODUCT(({due}<>"")*(DATE(YEAR(B{r}),MONTH(B{r})+(DAY(B{r})>{due}),{due})<F{r})*{amt}))', B.money)
        B.calc(ws.cell(r, 5), f'=IF(B{r}="","",C{r}-D{r})', B.money, True)
    ws.conditional_formatting.add(f'E11:E{10 + n}', CellIsRule(operator='lessThan', formula=['0'], font=Font(color=hexc(B.th['warn']), bold=True)))
    B.nav.append(('Paycheck', ws.title))
    ws.freeze_panes = 'A11'
    B.protect(ws)


# ---------------- build ----------------
SAMPLE_ROWS = [(3, 'Rent', 'Rent or mortgage', 1200), (4, 'Weekly shop', 'Groceries', 112.4), (6, 'Bus pass', 'Transport', 60),
               (9, 'Pizza night', 'Eating out', 38.5), (11, 'Weekly shop', 'Groceries', 96.2), (14, 'Phone bill', 'Phone and internet', 55),
               (15, 'To emergency fund', 'Emergency fund', 200), (18, 'Streaming', 'Subscriptions', 15.99), (21, 'Card payment', 'Credit card', 150), (24, 'Concert tickets', 'Fun', 120), (25, 'Shoes', 'Shopping', 64)]


def sample_data(B):
    """Fill January with an example month so buyers see it working (they clear it with one delete)."""
    if not getattr(B, 'month_sheets', None):
        return
    ws = B.wb[B.month_sheets[0]['title']]
    rows = SAMPLE_ROWS
    y = B.year or datetime.date.today().year + 1
    names = {c['name'] for c in B.cats}
    # typical rows in the extras, like the prefilled categories in Setup (buyers type over them)
    if getattr(B, 'bills', None):
        ws_b = B.wb[B.bills['title']]
        for i, (name, day, amt, auto) in enumerate([('Rent', 1, 1400, 'Yes'), ('Phone and internet', 14, 95, 'Yes'), ('Streaming', 18, 15.99, 'Yes'),
                                                     ('Car insurance', 20, 160, 'No'), ('Electric', 22, 90, 'No'), ('Credit card', 25, 150, 'No')]):
            r = B.bills['first'] + i
            ws_b.cell(r, 2).value = name; ws_b.cell(r, 3).value = day; ws_b.cell(r, 4).value = amt; ws_b.cell(r, 5).value = auto
    if 'debt' in B.wb.sheetnames or any(m == 'debt' for _, _, m in B.sheets):
        ws_d = B.wb[[t for _, t, m in B.sheets if m == 'debt'][0]]
        for i, (name, bal, rate, mn) in enumerate([('Credit card', 2400, 0.2299, 75), ('Car loan', 8200, 0.069, 210), ('Student loan', 14500, 0.055, 180)]):
            r = 10 + i
            ws_d.cell(r, 2).value = name; ws_d.cell(r, 3).value = bal; ws_d.cell(r, 4).value = rate; ws_d.cell(r, 5).value = mn
    if any(m == 'savings' for _, _, m in B.sheets):
        ws_s = B.wb[[t for _, t, m in B.sheets if m == 'savings'][0]]
        y0 = B.year or datetime.date.today().year + 1
        for i, (name, tgt, saved, when) in enumerate([('Emergency fund', 5000, 1200, datetime.date(y0, 12, 1)), ('Holiday trip', 1800, 450, datetime.date(y0, 7, 1)), ('New laptop', 1200, 300, datetime.date(y0, 9, 1))]):
            r = 7 + i
            ws_s.cell(r, 2).value = name; ws_s.cell(r, 3).value = tgt; ws_s.cell(r, 4).value = saved; ws_s.cell(r, 6).value = when
    for i, (d, what, cat, amt) in enumerate(rows):
        if cat not in names:
            continue
        r = 7 + i
        ws.cell(r, 2).value = datetime.date(y, 1, d); ws.cell(r, 3).value = what; ws.cell(r, 4).value = cat; ws.cell(r, 5).value = amt


def capabilities():
    return {'modules': {k: {'desc': v['desc'], 'params': v['params']} for k, v in MODULES.items()},
            'themes': list(THEMES), 'spec_fields': ['title', 'subtitle', 'brand', 'currency', 'year', 'theme', 'categories', 'income', 'sample_data', 'sections'],
            'notes': 'Delivered as .xlsx (opens in Excel and in Google Sheets via File > Import or a copy link). Fonts: use fonts that exist in both Google Sheets and Excel (Arial, Calibri fallback).'}


def validate_spec(spec):
    problems, needs = [], []
    for k in ('title', 'sections'):
        if k not in spec:
            problems.append(f'missing {k}')
    mods = [s.get('module') for s in spec.get('sections', [])]
    for m in mods:
        if m not in MODULES:
            needs.append(m)
    if ('months' in mods or 'dashboard' in mods) and 'setup' not in mods:
        problems.append('months and dashboard need a setup section first')
    if 'start_check' in mods and ('setup' not in mods or 'months' not in mods):
        problems.append('start_check needs setup and months sections')
    if 'paycheck' in mods and ('bills' not in mods or mods.index('paycheck') < mods.index('bills')):
        problems.append('paycheck must come after bills')
    if 'dashboard' in mods and ('months' not in mods or mods.index('dashboard') < mods.index('months')):
        problems.append('dashboard must come after months')
    return problems, needs


def build(spec, path):
    problems, needs = validate_spec(spec)
    if problems or needs:
        raise ValueError(f'spec problems: {problems}; needs_module: {needs}')
    B = Book(spec)
    for sec in spec['sections']:
        B.current = sec['module']
        MODULES[sec['module']]['fn'](B, sec)
    if spec.get('sample_data'):
        sample_data(B)
    for fn in B.after:
        fn()
    # put the dashboard tab right after Start here in the tab order
    order = [t for a, t, m in B.sheets]
    B.wb._sheets.sort(key=lambda ws: (0 if ws.title == order[0] else 1 if any(t == ws.title and m == 'dashboard' for a, t, m in B.sheets) else 2, order.index(ws.title)))
    B.save(path)
    return {'tabs': [t for t in B.wb.sheetnames], 'sheets': B.sheets, 'nav': B.nav}


if __name__ == '__main__':
    print(json.dumps(capabilities(), indent=1))
